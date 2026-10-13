"""Plain-language PowerPoint deck that explains the PDF report to non-technical readers."""
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from .plain import GLOSSARY

NAVY = RGBColor(0x12, 0x26, 0x3F)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
TEXT = RGBColor(0x27, 0x31, 0x3D)
MUTED = RGBColor(0x6B, 0x77, 0x85)
LIGHT = RGBColor(0xF3, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRADE = {"A": "1F9254", "B": "7CB342", "C": "E0A526", "D": "E07B2A", "F": "D64541", None: "9AA5B1"}
W, H = Emu(12192000), Emu(6858000)  # 16:9
IN = 914400


def _rgb(hexcode):
    return RGBColor.from_string(hexcode)


def _text(slide, x, y, w, h, text, size=18, bold=False, colour=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, colour
        run.font.name = "Calibri"
        p.space_after = Pt(6)
    return box


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def _slide(prs, title, kicker=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, 0.18, 7.5, TEAL)
    if kicker:
        _text(s, 0.6, 0.35, 11, 0.4, kicker.upper(), 12, True, TEAL)
    _text(s, 0.6, 0.7, 12, 0.9, title, 32, True, NAVY)
    return s


def _footer(prs, data):
    for i, s in enumerate(prs.slides, 1):
        if i == 1:
            continue
        _text(s, 0.6, 7.0, 9, 0.3, f"{data['legal']} · {data['target']} · {data['tier_label']}", 10, colour=MUTED)
        _text(s, 11.8, 7.0, 1.2, 0.3, str(i), 10, colour=MUTED, align=PP_ALIGN.RIGHT)


def _cards(slide, items, top=1.8, per_row=None):
    """items: list of (heading, body, colour_hex)."""
    n = len(items)
    per_row = per_row or min(n, 3)
    rows = (n + per_row - 1) // per_row
    gap = 0.25
    width = (12.1 - 0.6 - gap * (per_row - 1)) / per_row
    height = min(2.4, (6.7 - top - gap * (rows - 1)) / max(rows, 1))
    for i, (heading, body, colour) in enumerate(items):
        x = 0.6 + (i % per_row) * (width + gap)
        y = top + (i // per_row) * (height + gap)
        _rect(slide, x, y, width, height, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
        _rect(slide, x, y, 0.09, height, _rgb(colour))
        _text(slide, x + 0.25, y + 0.15, width - 0.4, 0.8, heading, 16, True, NAVY)
        _text(slide, x + 0.25, y + 0.85, width - 0.4, height - 0.9, body, 13, colour=TEXT)


SEV = {"critical": "D64541", "high": "E07B2A", "medium": "E0A526", "low": "6B7785"}
WHEN = {"critical": "Fix this week", "high": "Fix this month", "medium": "Plan within 3 months",
        "low": "When convenient"}


def render(data):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    full = data["tier"] == "comprehensive"
    score = data["score"]
    grade = score.get("grade")

    # 1. Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, 13.34, 7.5, NAVY)
    _rect(s, 0, 5.9, 13.34, 0.08, TEAL)
    _text(s, 0.8, 1.2, 11, 0.5, f"{data['brand']} CYBERHEALTH · {data['tier_label'].upper()}", 14, True,
          RGBColor(0x7F, 0xD6, 0xDC))
    _text(s, 0.8, 1.8, 11.5, 2.0, "Your security check, explained simply", 44, True, WHITE)
    _text(s, 0.8, 3.9, 11, 1.2, [data["client"], f"{data['target']} · {data['date']}"], 20,
          colour=RGBColor(0xC5, 0xD0, 0xDB))
    _text(s, 0.8, 6.2, 11.5, 0.6, "A plain-language companion to the full PDF report." +
          (" DRAFT: not yet reviewed." if data["draft"] else ""), 13, colour=RGBColor(0xC5, 0xD0, 0xDB))

    # 2. Score
    s = _slide(prs, "Your overall result", "The headline")
    circle = _rect(s, 0.9, 1.9, 3.4, 3.4, _rgb(GRADE[grade]), MSO_SHAPE.OVAL)
    circle.text_frame.text = ""
    _text(s, 0.9, 2.55, 3.4, 1.4, grade or "–", 88, True, WHITE, PP_ALIGN.CENTER)
    _text(s, 0.9, 4.1, 3.4, 0.6, f"{score.get('score', '–')} out of 100", 18, True, WHITE, PP_ALIGN.CENTER)
    _text(s, 4.9, 2.0, 7.6, 0.8, data["label"], 34, True, _rgb(GRADE[grade]))
    _text(s, 4.9, 2.9, 7.6, 2.5, [
        data["interpretation"],
        f"We found {len(data['findings'])} thing(s) to improve and {len(data['positives'])} thing(s) already working well.",
        f"This result covers {score.get('coverage', 0)}% of our full check." +
        (" Some areas still need to be assessed." if data["not_assessed"] else ""),
    ], 18)
    _text(s, 0.9, 5.8, 11.5, 0.6, "Grades: A Excellent · B Good · C Moderate · D Needs improvement · F Critical", 13,
          colour=MUTED)

    # 3. Areas (comprehensive)
    if full:
        s = _slide(prs, "How each area did", "Score by area")
        for i, a in enumerate(data["areas"]):
            y = 1.85 + i * 0.6
            _text(s, 0.6, y, 4.4, 0.5, a["plain"], 15, True, NAVY)
            _rect(s, 5.1, y + 0.1, 6.0, 0.32, LIGHT)
            if a["score"] is not None:
                _rect(s, 5.1, y + 0.1, max(0.05, 6.0 * a["score"] / 100), 0.32, _rgb(GRADE[a["grade"]]))
            _text(s, 11.3, y, 1.6, 0.5, f"{a['score']}" if a["score"] is not None else "Not checked", 15, True,
                  NAVY if a["score"] is not None else MUTED)

    # 4. Top fixes
    top = data["findings"][: (6 if full else 3)]
    s = _slide(prs, "The most important things to fix", "Priorities")
    if top:
        _cards(s, [(f["title"], f"Why it matters: {f['impact']}\nWhat to do: {f['action']}\nWhen: {WHEN[f['severity']]}",
                    SEV[f["severity"]]) for f in top], per_row=3)
    else:
        _text(s, 0.6, 2.0, 11, 1, "Nothing urgent was found. Keep the current good practices in place.", 22)

    # 5. What's working
    s = _slide(prs, "What is already working well", "Strengths")
    good = [p["good"] for p in data["positives"][: (8 if full else 5)]]
    _text(s, 0.6, 1.8, 12, 5, ["✔  " + g for g in good] or ["No passing controls in the assessed areas yet."], 18)

    # 6. Plan (comprehensive)
    if full:
        s = _slide(prs, "Your action plan", "Roadmap")
        cols = [("This week", [f for f in data["findings"] if f["severity"] == "critical"], "D64541"),
                ("This month", [f for f in data["findings"] if f["severity"] == "high"], "E07B2A"),
                ("Next 3 months", [f for f in data["findings"] if f["severity"] in ("medium", "low")], "E0A526")]
        for i, (name, items, colour) in enumerate(cols):
            x = 0.6 + i * 4.1
            _rect(s, x, 1.8, 3.85, 0.6, _rgb(colour))
            _text(s, x + 0.2, 1.88, 3.5, 0.5, name, 18, True, WHITE)
            _rect(s, x, 2.4, 3.85, 4.3, LIGHT)
            lines = [f"• {f['action']}" for f in items[:6]] or ["Nothing in this period."]
            if len(items) > 6:
                lines.append(f"… and {len(items) - 6} more in the PDF report")
            _text(s, x + 0.2, 2.55, 3.5, 4.0, lines, 13)

    # 7. Next steps
    s = _slide(prs, "Next steps", "What happens now")
    steps = ["Share this summary with whoever looks after your website, email and hosting.",
             "Start with the items marked 'Fix this week' and 'Fix this month'."]
    if data["not_assessed"]:
        steps.append("Complete the areas not yet checked: " + ", ".join(a["plain"].lower() for a in data["not_assessed"]) + ".")
    steps.append("Book a follow-up check after the fixes to confirm they worked." if full else
                 "Upgrade to a Comprehensive Assessment for the full findings list, a detailed action plan and deeper checks.")
    _text(s, 0.6, 1.8, 12, 4.5, [f"{i}.  {t}" for i, t in enumerate(steps, 1)], 20)

    # 8. Glossary (comprehensive)
    if full:
        s = _slide(prs, "Words you may hear", "Glossary")
        _text(s, 0.6, 1.8, 12, 5, [f"{term}: {meaning}" for term, meaning in GLOSSARY[:9]], 15)

    _footer(prs, data)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
